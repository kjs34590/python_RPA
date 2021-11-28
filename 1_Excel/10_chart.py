from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, LineChart

wb = load_workbook("./PYTHON_WORKS/python_RPA/1_Excel/cell_range.xlsx")
ws = wb.active

# B2:C11까지의 데이터를 Bar 차트로 생성
bar_value = Reference(ws, min_row=2, max_row=11, min_col=2, max_col=3)
bar_chart = BarChart() # 차트 종류 설정 (Bar, Line, Pie, ...)
bar_chart.add_data(bar_value) # 차트 데이터 추가

ws.add_chart(bar_chart, "E1") # 차트 넣을 위치 정의

# B1:C11까지의 데이터를 Line 차트로 생성. (B1을 포함해 계열까지 가져온다.)
line_value = Reference(ws, min_row=1, max_row=11, min_col=2, max_col=3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data = True) # True를 주면 첫 row 데이터를 계열 값으로 쓴다.
line_chart.title = "영어/수학 성적"
line_chart.style = 10 # 미리 정의된 스타일 적용 (사용자가 개별 지정도 가능)
line_chart.y_axis.title = "점수" # Y축 제목
line_chart.x_axis.title = "번호"
ws.add_chart(line_chart, "E15")

wb.save("./PYTHON_WORKS/python_RPA/1_Excel/sample_chart.xlsx")